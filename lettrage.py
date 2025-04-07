import streamlit as st
from io import BytesIO
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

# Cores para preenchimento
verde = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
azul = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")

st.title("Lettrage de Excel - Versão Web")

uploaded_file = st.file_uploader("Escolhe um ficheiro Excel", type=["xlsx"])

if uploaded_file:
    st.success("Ficheiro carregado com sucesso.")

    if st.button("Processar ficheiro"):
        try:
            # Carregar o Excel
            in_memory_file = BytesIO(uploaded_file.read())
            wb = load_workbook(in_memory_file)
            folha_original = wb.worksheets[0]
            folha_valores = wb.copy_worksheet(folha_original)
            folha_valores.title = "Pares_Coluna_J"
            folha_ultimos8 = wb.copy_worksheet(folha_original)
            folha_ultimos8.title = "Ultimos8_L_em_ELM"

            max_linhas = folha_original.max_row
            max_colunas = folha_original.max_column

            # --- Parte 1: Pares positivos/negativos na coluna J ---
            valores_j = {}
            for i in range(2, max_linhas + 1):
                valor = folha_valores.cell(row=i, column=10).value
                if isinstance(valor, (int, float)) and abs(valor) >= 0.01:
                    valores_j[i] = valor

            usados = set()
            for i, val1 in valores_j.items():
                if i in usados:
                    continue
                for j, val2 in valores_j.items():
                    if j == i or j in usados:
                        continue
                    if abs(val1 + val2) <= 0.01:
                        usados.update([i, j])
                        folha_valores.cell(row=i, column=21).value = f"Linha {j}"
                        folha_valores.cell(row=j, column=21).value = f"Linha {i}"
                        for c in range(1, max_colunas + 1):
                            folha_valores.cell(row=i, column=c).fill = verde
                            folha_valores.cell(row=j, column=c).fill = verde
                        break

            # --- Parte 2: Últimos 8 caracteres da coluna L em E, L, M (outras linhas) ---
            for i in range(2, max_linhas + 1):
                celula_L = folha_ultimos8.cell(row=i, column=12).value  # Coluna L
                if celula_L is None:
                    continue

                ult8 = str(celula_L)[-8:]
                linhas_encontradas = []

                for j in range(2, max_linhas + 1):
                    if j == i:
                        continue
                    val_E = str(folha_ultimos8.cell(row=j, column=5).value or "")
                    val_L = str(folha_ultimos8.cell(row=j, column=12).value or "")
                    val_M = str(folha_ultimos8.cell(row=j, column=13).value or "")

                    if ult8 in val_E or ult8 in val_L or ult8 in val_M:
                        linhas_encontradas.append(j)
                        for c in range(1, max_colunas + 1):
                            folha_ultimos8.cell(row=j, column=c).fill = azul

                if linhas_encontradas:
                    folha_ultimos8.cell(row=i, column=21).value = f"Linhas: {', '.join(map(str, linhas_encontradas))}"

            # Guardar em memória e disponibilizar para download
            output = BytesIO()
            wb.save(output)
            output.seek(0)

            st.success("Processamento concluído!")
            st.download_button(
                label="📥 Descarregar ficheiro processado",
                data=output,
                file_name="ficheiro_processado.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

        except Exception as e:
            st.error(f"Ocorreu um erro: {e}")